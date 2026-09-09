"""Read the Excel scaffold: cached outputs for reconciliation and base inputs for drift checks.

The workbook is the author's scaffold, not evidence. This module never writes it.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .provenance import PACKAGE_ROOT
from .schemas import FROZEN_STRATEGY_IDS, StrategyId

WORKBOOK_PATH = (
    PACKAGE_ROOT
    / "data"
    / "raw_snapshots"
    / "workbook_scaffold"
    / "2026-09-01"
    / "Telo_Feasibility_Model.xlsx"
)
WORKBOOK_SHA256 = "367069513203098f0157dfca83272e7bda35b045ecd67139bb268d483b9be9f8"

# Workbook row order (09_Deterministic rows 5-12 = product A, 13-20 = product B).
WORKBOOK_PRODUCT_ORDER: tuple[str, str] = (
    "sodium_bicarbonate_8_4_50ml",
    "norepinephrine_1mgml_4ml",
)
WORKBOOK_PRODUCT_LABEL: dict[str, str] = {
    "sodium_bicarbonate_8_4_50ml": "Product A",
    "norepinephrine_1mgml_4ml": "Product B",
}
WORKBOOK_STRATEGY_NAME: dict[StrategyId, str] = {
    StrategyId.S0: "Status quo centralized",
    StrategyId.S1: "Additional safety stock",
    StrategyId.S2: "Dual source + contracts",
    StrategyId.S3: "Reserved contract capacity",
    StrategyId.S4: "Additional centralized capacity",
    StrategyId.S5: "Distributed regional nodes",
    StrategyId.S6: "Distributed + validated release assurance",
    StrategyId.S7: "503B shortage response",
}
# Typed eligibility flags in 08_Strategies!L5:L12. A workbook artifact (WB-05): used ONLY by the
# replica so that 12_Results reconciles; the corrected variant derives eligibility from gates.
WORKBOOK_ELIGIBILITY_FLAG: dict[StrategyId, str] = {
    StrategyId.S0: "YES",
    StrategyId.S1: "YES",
    StrategyId.S2: "YES",
    StrategyId.S3: "YES",
    StrategyId.S4: "UNCERTAIN",
    StrategyId.S5: "UNCERTAIN",
    StrategyId.S6: "UNCERTAIN",
    StrategyId.S7: "UNCERTAIN",
}

SCREEN_COLUMNS: dict[str, str] = {
    "C": "sites",
    "D": "annual_demand",
    "E": "units_per_batch",
    "F": "batches_per_site_year",
    "G": "yield_fraction",
    "H": "uptime",
    "I": "saleable_capacity",
    "J": "utilization",
    "K": "annualized_capex",
    "L": "fixed_qa_labor",
    "M": "validation_annualized",
    "N": "variable_production",
    "O": "testing",
    "P": "inventory_carrying_expiry",
    "Q": "distribution",
    "R": "reserved_capacity_cost",
    "S": "total_annual_cost",
    "T": "cost_per_unit",
    "U": "response_days",
    "V": "expected_shortage_days_screen",
}
DASHBOARD_COLUMNS: dict[str, str] = {
    "C": "cost_per_unit",
    "D": "response_days",
    "E": "expected_shortage_days",
    "F": "shortage_days_avoided",
    "G": "incremental_annual_cost",
    "H": "cost_per_shortage_day_avoided",
    "I": "approx_fill_rate",
    "J": "interpretation",
}

# Product-sheet row -> parameter id (06_Product_A / 07_Product_B rows 5-24).
PRODUCT_ROWS: dict[int, str] = {
    5: "annual_demand_units",
    6: "units_per_batch",
    7: "batches_per_site_year_nominal",
    8: "yield_fraction",
    9: "uptime_fraction",
    10: "production_cycle_days",
    11: "release_time_days",
    12: "material_lead_time_days",
    13: "changeover_days",
    14: "shelf_life_months",
    15: "variable_materials_usd_per_unit",
    16: "variable_conversion_usd_per_unit",
    17: "testing_usd_per_batch",
    18: "fixed_qa_labor_usd_per_site_year",
    19: "capital_usd_per_site",
    20: "validation_usd_one_time",
    21: "distribution_usd_per_unit",
    22: "delivery_days",
    23: "expiry_scrap_fraction",
    24: "target_safety_stock_days",
}
GLOBAL_ROWS: dict[int, str] = {
    6: "discount_rate",
    7: "capital_economic_life_years",
    8: "inventory_carrying_rate",
    14: "annual_demand_growth",
    15: "demand_cv",
    16: "demand_shock_multiplier",
    17: "demand_shocks_per_year",
    18: "demand_shock_duration_days",
    19: "site_failures_per_site_year",
    20: "site_failure_duration_days",
    21: "common_cause_events_per_year",
    22: "common_cause_duration_days",
    23: "common_cause_capacity_impact",
    24: "unmet_demand_severity",
    25: "capital_recovery_factor_typed",
}
STRATEGY_COLUMNS: dict[str, str] = {
    "B": "sites",
    "C": "capacity_factor",
    "D": "safety_stock_days",
    "E": "release_time_factor",
    "F": "delivery_days",
    "G": "common_impact_factor",
    "H": "fixed_cost_factor",
    "I": "validation_factor",
    "J": "reserved_capacity_fraction",
}


@dataclass(frozen=True)
class WorkbookRow:
    sheet: str
    row: int
    product_id: str
    strategy_id: StrategyId
    values: dict[str, Any]


def _row_identity(row: int) -> tuple[str, StrategyId]:
    if not 5 <= row <= 20:
        raise ValueError(f"row {row} outside 5..20")
    product = WORKBOOK_PRODUCT_ORDER[0] if row <= 12 else WORKBOOK_PRODUCT_ORDER[1]
    strategy = FROZEN_STRATEGY_IDS[(row - 5) % 8]
    return product, strategy


def cached_screen_rows(path: Path = WORKBOOK_PATH) -> list[WorkbookRow]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["09_Deterministic"]
    rows: list[WorkbookRow] = []
    for r in range(5, 21):
        product, strategy = _row_identity(r)
        vals = {name: ws[f"{col}{r}"].value for col, name in SCREEN_COLUMNS.items()}
        rows.append(WorkbookRow("09_Deterministic", r, product, strategy, vals))
    wb.close()
    return rows


def cached_dashboard_rows(path: Path = WORKBOOK_PATH) -> list[WorkbookRow]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["12_Results"]
    rows: list[WorkbookRow] = []
    for r in range(5, 21):
        product, strategy = _row_identity(r)
        vals = {name: ws[f"{col}{r}"].value for col, name in DASHBOARD_COLUMNS.items()}
        rows.append(WorkbookRow("12_Results", r, product, strategy, vals))
    wb.close()
    return rows


@dataclass(frozen=True)
class WorkbookInputs:
    global_base: dict[str, float]
    global_low: dict[str, float]
    global_high: dict[str, float]
    product_base: dict[str, dict[str, float]]
    product_low: dict[str, dict[str, float]]
    product_high: dict[str, dict[str, float]]
    strategies: dict[StrategyId, dict[str, float]]
    eligibility_flags: dict[StrategyId, str]
    target_fill_rate: float


def workbook_inputs(path: Path = WORKBOOK_PATH) -> WorkbookInputs:
    """Base/low/high inputs as typed in the workbook (05, 06, 07, 08). Drift guard for config/."""
    wb = load_workbook(path, data_only=True, read_only=True)
    g = wb["05_Assumptions"]
    gb = {name: float(g[f"C{r}"].value) for r, name in GLOBAL_ROWS.items()}
    gl = {name: float(g[f"B{r}"].value) for r, name in GLOBAL_ROWS.items()}
    gh = {name: float(g[f"D{r}"].value) for r, name in GLOBAL_ROWS.items()}
    target_fill = float(g["C9"].value)
    pb: dict[str, dict[str, float]] = {}
    pl: dict[str, dict[str, float]] = {}
    ph: dict[str, dict[str, float]] = {}
    for pid, sheet in zip(WORKBOOK_PRODUCT_ORDER, ("06_Product_A", "07_Product_B"), strict=True):
        ws = wb[sheet]
        pb[pid] = {name: float(ws[f"C{r}"].value) for r, name in PRODUCT_ROWS.items()}
        pl[pid] = {name: float(ws[f"B{r}"].value) for r, name in PRODUCT_ROWS.items()}
        ph[pid] = {name: float(ws[f"D{r}"].value) for r, name in PRODUCT_ROWS.items()}
    s = wb["08_Strategies"]
    strategies: dict[StrategyId, dict[str, float]] = {}
    flags: dict[StrategyId, str] = {}
    for i, sid in enumerate(FROZEN_STRATEGY_IDS):
        r = 5 + i
        strategies[sid] = {
            name: float(s[f"{col}{r}"].value) for col, name in STRATEGY_COLUMNS.items()
        }
        flags[sid] = str(s[f"L{r}"].value)
    wb.close()
    return WorkbookInputs(gb, gl, gh, pb, pl, ph, strategies, flags, target_fill)
